"""Experimental export per LLC plus the inter-entity statement (decisions 13, 14, 26).

One Excel file per LLC and one inter-entity file, zipped. Every sheet carries
an EXPERIMENTAL banner: these are prototype results, not accounting entries.
"""
from __future__ import annotations

import io
import zipfile
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

BANNER = "EXPERIMENTAL — C-Bank Prototype dry run. Not accounting entries; not approved; Bank data unchanged."
COLUMNS = [("txn_id", "Txn"), ("date", "Date"), ("description", "Description"), ("amount", "Amount"),
           ("source", "Account/Card"), ("payer", "Paid by"), ("beneficiary", "Beneficiary"), ("who_name", "WHO"),
           ("why", "WHY"), ("category", "Category RF-One"), ("qb", "QB account"), ("operation_type", "Operation"),
           ("status", "Status"), ("methods", "Method WHO/WHY/Beneficiary"), ("protected", "Human decision"),
           ("comparison", "Compared with Bank"), ("issues", "Exceptions / anomalies")]


def _row(r: dict) -> list:
    values = dict(r)
    values["amount"] = r["amount_minor"] / 100
    values["qb"] = r["qb_account"] or "not mapped"
    values["methods"] = f"{r['who_method'] or '—'} / {r['why_method'] or '—'} / {r['beneficiary_method'] or '—'}"
    values["issues"] = "; ".join(f"{e['code']}: {e['message']}" for e in r["exceptions"] + r["anomalies"])
    return [values.get(k) for k, _ in COLUMNS]


def _sheet(wb, title, rows, columns=COLUMNS, convert=_row):
    ws = wb.create_sheet(title[:31])
    ws.append([BANNER])
    ws["A1"].font = Font(bold=True, color="9C0006")
    ws["A1"].fill = PatternFill("solid", fgColor="FFC7CE")
    ws.append([label for _, label in columns])
    for cell in ws[2]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(convert(r))
    ws.freeze_panes = "A3"
    return ws


def balances(results: list[dict]) -> dict[tuple, int]:
    """Net amounts owed: (debtor, creditor) -> minor units (decision 14)."""
    net = defaultdict(int)
    for r in results:
        led = r.get("ledger")
        if not led or not led.get("beneficiary"):
            continue
        if led["kind"] == "PAID_FOR_OTHER":
            debtor, creditor = (led["beneficiary"], led["payer"]) if led["amount_minor"] > 0 else (led["payer"], led["beneficiary"])
            net[(debtor, creditor)] += abs(led["amount_minor"])
        elif led["kind"] == "TRANSFER":
            # Money sent from payer to beneficiary: reduces what payer owes, or creates an advance.
            net[(led["beneficiary"], led["payer"])] += led["amount_minor"]
    out = {}
    for (a, b), v in net.items():
        if (b, a) in out:
            continue
        diff = v - net.get((b, a), 0)
        if diff > 0:
            out[(a, b)] = diff
        elif diff < 0:
            out[(b, a)] = -diff
    return out


def build_zip(results: list[dict], entities: dict, run_label: str) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for e in entities.values():
            if e.entity_type != "LEGAL":
                continue
            wb = Workbook()
            wb.remove(wb.active)
            attributed = [r for r in results if r["beneficiary_id"] == e.id]
            paid_for_others = [r for r in results if r["payer_id"] == e.id and r["beneficiary_id"] not in (None, e.id)
                               and r.get("ledger")]
            open_items = [r for r in results if (r["beneficiary_id"] == e.id or r["payer_id"] == e.id)
                          and (r["status"] != "RECONCILED" or r["anomalies"])]
            differences = [r for r in results if (r["beneficiary_id"] == e.id or r["payer_id"] == e.id)
                           and (r["comparison"] or "").startswith("DIFFERENT")]
            _sheet(wb, "Transactions", attributed)
            _sheet(wb, "Paid for others", paid_for_others)
            _sheet(wb, "Anomalies", open_items)
            _sheet(wb, "Differences vs Bank", differences)
            out = io.BytesIO()
            wb.save(out)
            z.writestr(f"CBank_EXPERIMENTAL_{e.code}_{run_label}.xlsx", out.getvalue())
        wb = Workbook()
        wb.remove(wb.active)
        ledger_rows = [r for r in results if r.get("ledger")]
        lcols = [("txn_id", "Txn"), ("date", "Date"), ("description", "Description"), ("kind", "Kind"),
                 ("payer", "Paid by"), ("beneficiary", "Beneficiary"), ("amount", "Amount"),
                 ("proposal", "Proposal"), ("note", "Note")]
        _sheet(wb, "Inter-entity movements", ledger_rows, lcols,
               lambda r: [r["txn_id"], r["date"], r["description"], r["ledger"]["kind"], r["ledger"]["payer"],
                          r["ledger"]["beneficiary"], r["ledger"]["amount_minor"] / 100, r["ledger"]["proposal"],
                          r["ledger"]["note"]])
        bal = [dict(debtor=a, creditor=b, amount=v / 100) for (a, b), v in sorted(balances(results).items())]
        _sheet(wb, "Proposed net balances", bal, [("debtor", "Owes"), ("creditor", "To"), ("amount", "Net amount")],
               lambda r: [r["debtor"], r["creditor"], r["amount"]])
        out = io.BytesIO()
        wb.save(out)
        z.writestr(f"CBank_EXPERIMENTAL_InterEntity_{run_label}.xlsx", out.getvalue())
    return buf.getvalue()
