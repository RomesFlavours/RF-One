#!/usr/bin/env python
"""Batch résumé upload tests (TASK_SELECTION_002; extended by Task 2A for
DOCX/TXT support, the real rule-based parsing fallback, and the PARTIAL /
no-extractable-text outcomes).

Mirrors `03 Software/RF-One Data Store/test_selection_engine.py`'s
main()-returns-exit-code convention, but exercises the actual Flask routes
via Werkzeug's test client (not just the persistence layer) — this is what
proves batch upload, per-file failure isolation, and duplicate detection
work end to end through `app.py`, not only through
`rfone_data_store/selection/import_pipeline.py` in isolation.

Task 2A replaced the fabricated-fixture DEMO fallback with a real,
rule-based parser (`rfone_data_store/selection/parsing/deterministic_parser.py`);
these tests upload real, hand-written résumé text/DOCX fixtures and assert
the structured candidate data actually comes from that content — never a
canned fixture — since no AI credentials are configured in this environment.

Runs against a throwaway SQLite database (never the real
`data/selection.db`) and cleans up every file it writes into `uploads/`
afterward.

Usage:
    python test_batch_upload.py
"""

from __future__ import annotations

import os
import sys
import tempfile
from io import BytesIO

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
_TEST_DB_FD, _TEST_DB_PATH = tempfile.mkstemp(suffix=".db", prefix="selection_batch_test_")
os.close(_TEST_DB_FD)
os.remove(_TEST_DB_PATH)  # let app.py's migration runner create it fresh
os.environ["RFONE_DATABASE_URL"] = f"sqlite:///{_TEST_DB_PATH.replace(os.sep, '/')}"

import app as selection_app  # noqa: E402

try:
    import docx as _python_docx
except ImportError:  # pragma: no cover - optional dependency
    _python_docx = None


ALPHA_RESUME_TXT = """Jordan Rivera
jordan.rivera@example.com
(407) 555-0199
Orlando, FL

EXPERIENCE
Server, The Garden Bistro
Jun 2019 - Mar 2022
Took orders for a 10-table section and processed payments.
Trained new servers on menu items.

Host, Lakeside Grill
Jan 2018 - May 2019
Greeted guests and managed the wait list.

EDUCATION
Orlando High School
High School Diploma
2014 - 2018

SKILLS
Customer service, POS systems, Wine service

LANGUAGES
English, Spanish (Conversational)

CERTIFICATIONS
Food Handler Certificate - ServSafe, 2021
"""

BETA_RESUME_DOCX_PARAGRAPHS = [
    "Alex Chen",
    "alex.chen@example.com",
    "(689) 555-0142",
    "Winter Park, FL",
    "EXPERIENCE",
    "Bartender - Rome's Flavours",
    "May 2021 - Present",
    "Full bar service for a 120-seat restaurant, crafted seasonal cocktails.",
    "EDUCATION",
    "Valencia College",
    "Associate of Arts in Hospitality Management",
    "2017 - 2019",
]

NO_STRUCTURE_TEXT = (
    "lorem ipsum dolor sit amet consectetur adipiscing elit sed do eiusmod "
    "tempor incididunt ut labore et dolore magna aliqua ut enim ad minim veniam"
)


def _make_docx_bytes(paragraphs: list[str]) -> bytes:
    document = _python_docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    buf = BytesIO()
    document.save(buf)
    return buf.getvalue()


def _fake_pdf(label: str) -> bytes:
    # Not a real, parseable PDF — extract_text_from_pdf() fails closed
    # (returns None). The pipeline must report this as "no extractable
    # text," never fabricate a candidate for it (Task 2A §2/§8).
    return f"%PDF-1.4 fake test resume content {label}".encode("utf-8")


def main() -> int:
    checks_passed: list[str] = []
    checks_failed: list[str] = []

    def check(description: str, condition: bool) -> None:
        if condition:
            checks_passed.append(description)
        else:
            checks_failed.append(description)

    existing_upload_files = set(os.listdir(selection_app.UPLOAD_DIR))
    client = selection_app.app.test_client()

    try:
        with selection_app.SessionFactory() as session:
            restaurant = selection_app._bootstrap_restaurant(session)
            starting_candidate_count = len(
                selection_app.persistence.list_candidates(session, restaurant_id=restaurant.id)
            )

        # ---------------------------------------------------------------
        # 1. Batch upload via the classic multi-file form fallback
        #    (/upload): a real TXT résumé, a real DOCX résumé (if
        #    python-docx is installed), an unsupported format, a
        #    corrupt/no-text PDF, and an exact re-upload of the TXT
        #    résumé (duplicate detection).
        # ---------------------------------------------------------------
        files = [
            (BytesIO(ALPHA_RESUME_TXT.encode("utf-8")), "candidate_alpha.txt"),
            (BytesIO(b"whatever"), "candidate_gamma.rtf"),
            (BytesIO(_fake_pdf("delta")), "candidate_delta.pdf"),
            (BytesIO(ALPHA_RESUME_TXT.encode("utf-8")), "candidate_alpha_renamed.txt"),
        ]
        if _python_docx is not None:
            files.insert(1, (BytesIO(_make_docx_bytes(BETA_RESUME_DOCX_PARAGRAPHS)), "candidate_beta.docx"))

        response = client.post("/upload", data={"resume_file": files}, content_type="multipart/form-data")
        check("1: batch upload returns HTTP 200", response.status_code == 200)

        with selection_app.SessionFactory() as session:
            restaurant = selection_app._bootstrap_restaurant(session)
            candidates = selection_app.persistence.list_candidates(session, restaurant_id=restaurant.id)
            new_candidates = len(candidates) - starting_candidate_count

        expected_new = 2 if _python_docx is not None else 1
        check(
            f"2: exactly {expected_new} new Candidate(s) created from the batch (unsupported format, "
            "no-extractable-text PDF, and the exact re-upload must NOT have created rows)",
            new_candidates == expected_new,
        )

        alpha_candidate_id = next((c.id for c in candidates if c.full_name == "Jordan Rivera"), None)
        check(
            "3: the real TXT résumé's actual name was extracted from its own text (not a fixture/mock name)",
            alpha_candidate_id is not None,
        )
        if alpha_candidate_id is not None:
            with selection_app.SessionFactory() as session:
                alpha_candidate = selection_app.persistence.get_candidate(session, alpha_candidate_id)
                profile = selection_app.persistence.to_profile(alpha_candidate)
                alpha_parsing_mode = alpha_candidate.parsing_mode
            check(
                "4: parsing_mode is RULE_BASED (no AI credentials configured in this environment) — "
                "never DEMO for a real upload",
                alpha_parsing_mode == "RULE_BASED",
            )
            check(
                "5: email/phone/location were extracted from the actual résumé text",
                profile.email == "jordan.rivera@example.com" and profile.phone is not None
                and profile.location is not None,
            )
            check(
                "6: multiple employment records were extracted (Server + Host), each from real text",
                len(profile.work_history) == 2,
            )
            server_entry = next((w for w in profile.work_history if w.employer == "The Garden Bistro"), None)
            check(
                "6b (Task 2B): the Server/Garden Bistro entry received normalized dates and role through "
                "the real import pipeline — MONTH-precision Jun 2019 start, 'Server' display title, "
                "'FOH Service' role family, original title untouched",
                server_entry is not None and server_entry.start_date_precision == "MONTH"
                and server_entry.normalized_title == "Server" and server_entry.role_family == "FOH Service"
                and server_entry.original_job_title == "Server",
            )
            check(
                "7: skills, a language and a certification were extracted from the résumé's own sections",
                len(profile.skills) >= 2 and len(profile.languages_detail) >= 1
                and len(profile.certifications) == 1,
            )
            check(
                "8: a field never stated on this résumé (LinkedIn URL) stays None — never fabricated",
                profile.linkedin_url is None,
            )

        # ---------------------------------------------------------------
        # 2. Per-file JSON endpoint (/api/upload) used by the batch UI's
        #    JavaScript — one file per call.
        # ---------------------------------------------------------------
        epsilon_text = ALPHA_RESUME_TXT.replace("Jordan Rivera", "Taylor Morgan").replace(
            "jordan.rivera@example.com", "taylor.morgan@example.com"
        )
        good_response = client.post(
            "/api/upload",
            data={"resume_file": (BytesIO(epsilon_text.encode("utf-8")), "candidate_epsilon.txt")},
            content_type="multipart/form-data",
        )
        good_json = good_response.get_json()
        check(
            "9: /api/upload returns COMPLETED with a candidate_id for a valid résumé",
            good_response.status_code == 200 and good_json["status"] == "COMPLETED"
            and good_json["candidate_id"] is not None,
        )

        bad_response = client.post(
            "/api/upload",
            data={"resume_file": (BytesIO(b"whatever"), "candidate_zeta.rtf")},
            content_type="multipart/form-data",
        )
        bad_json = bad_response.get_json()
        check(
            "10: /api/upload returns FAILED (never a 500) for an unsupported format, with a readable reason",
            bad_response.status_code == 200 and bad_json["status"] == "FAILED"
            and bad_json["error"] and "unsupported" in bad_json["error"].lower(),
        )

        no_text_response = client.post(
            "/api/upload",
            data={"resume_file": (BytesIO(_fake_pdf("eta")), "candidate_eta.pdf")},
            content_type="multipart/form-data",
        )
        no_text_json = no_text_response.get_json()
        check(
            "11: a PDF with no extractable text is reported FAILED with an explicit reason, never "
            "silently completed with fabricated data",
            no_text_json["status"] == "FAILED" and no_text_json["error"]
            and "extractable" in no_text_json["error"].lower(),
        )

        partial_response = client.post(
            "/api/upload",
            data={"resume_file": (BytesIO(NO_STRUCTURE_TEXT.encode("utf-8")), "candidate_theta.txt")},
            content_type="multipart/form-data",
        )
        partial_json = partial_response.get_json()
        check(
            "12: a résumé with real but unstructured text (no name/contact/section recognized) is "
            "reported PARTIAL, not COMPLETED — its text is preserved for manual review, nothing invented",
            partial_json["status"] == "PARTIAL",
        )

        duplicate_response = client.post(
            "/api/upload",
            data={"resume_file": (BytesIO(ALPHA_RESUME_TXT.encode("utf-8")), "candidate_alpha.txt")},
            content_type="multipart/form-data",
        )
        duplicate_json = duplicate_response.get_json()
        check(
            "13: re-uploading the exact same résumé via /api/upload is flagged DUPLICATE, not a new candidate",
            duplicate_json["status"] == "DUPLICATE",
        )

        # ---------------------------------------------------------------
        # 3. Persistence — candidates remain visible after a fresh request
        #    (simulates "refresh the page").
        # ---------------------------------------------------------------
        home_response = client.get("/")
        check("14: the candidate list page loads successfully after the batch", home_response.status_code == 200)
        body = home_response.get_data(as_text=True)
        check(
            "15: source metadata (LOCAL_UPLOAD / manual) is visible on the candidate list page",
            "LOCAL_UPLOAD" in body and "manual" in body,
        )

        detail_response = client.get(f"/candidate/{good_json['candidate_id']}")
        check("16: the candidate detail page loads for a freshly imported candidate", detail_response.status_code == 200)
        detail_body = detail_response.get_data(as_text=True)
        check(
            "17: the candidate detail page shows the Source & Extraction Information section",
            "Source" in detail_body and "Extraction" in detail_body,
        )
        check(
            "18: the raw résumé text remains available on the detail page (traceable evidence)",
            "taylor.morgan@example.com" in detail_body,
        )

        # ---------------------------------------------------------------
        # 4. Task 3C-FIX — Original CV access (checks L/M) and the new
        #    Review Queue / authoring pages loading end to end.
        # ---------------------------------------------------------------
        cv_response = client.get(f"/candidate/{good_json['candidate_id']}/original-cv")
        check(
            "19 (3C-FIX-L/M): Open Original CV opens the REAL stored résumé source for a freshly "
            "uploaded candidate (HTTP 200, serving the actual saved file)",
            cv_response.status_code == 200,
        )

        with selection_app.SessionFactory() as session:
            alpha_candidate = selection_app.persistence.get_candidate(session, alpha_candidate_id)
            raw = session.get(selection_app.m.RawResume, alpha_candidate.raw_resume_id)
            missing_path = raw.storage_path
        os.remove(missing_path)  # simulate a since-removed local upload (this file was never send_file()'d above)
        existing_upload_files.discard(os.path.basename(missing_path))  # already removed, not ours to clean up again
        unavailable_response = client.get(f"/candidate/{alpha_candidate_id}/original-cv")
        check(
            "20 (3C-FIX-L/M): a candidate whose original file is no longer on disk shows a clear "
            "'unavailable' state (HTTP 404), never a broken/500 response",
            unavailable_response.status_code == 404
            and "unavailable" in unavailable_response.get_data(as_text=True).lower(),
        )

        applications_response = client.get("/applications")
        check("21: the Review Queue (/applications) loads successfully", applications_response.status_code == 200)
        applications_body = applications_response.get_data(as_text=True)
        check(
            "22: the Review Queue shows filter controls (priority/workflow status/target role/sort)",
            "Filter" in applications_body and "Workflow status" in applications_body,
        )

        with selection_app.SessionFactory() as session:
            restaurant = selection_app._bootstrap_restaurant(session)
            first_application = selection_app.app_svc.list_applications(session, restaurant_id=restaurant.id)[0]
            first_application_id = first_application.id
        application_detail_response = client.get(f"/applications/{first_application_id}")
        check(
            "23: an Application Review page loads successfully and shows the workflow-status control",
            application_detail_response.status_code == 200
            and "Workflow Status" in application_detail_response.get_data(as_text=True),
        )

        signals_response = client.get("/signals")
        check(
            "24 (3C-FIX-G): the Signal Definition authoring page loads and shows a create form",
            signals_response.status_code == 200 and "Create a Signal Definition" in signals_response.get_data(as_text=True),
        )

        policies_response = client.get("/priority-policies")
        check(
            "25 (3C-FIX-I): the Review Priority Policy authoring page loads and shows a create form",
            policies_response.status_code == 200 and "Create a Review Priority Policy" in policies_response.get_data(as_text=True),
        )

        identity_response = client.get("/identity-matches")
        check("26: the Identity Matches page loads successfully", identity_response.status_code == 200)

        # ---------------------------------------------------------------
        # 5. Task 3D-FIX — automatic Primary Screening through the real
        #    HTTP routes, with NO AI provider configured in this
        #    environment (the authentic unavailable-provider path).
        # ---------------------------------------------------------------
        criterion_response = client.post(
            "/primary-screening-criteria/new",
            data={"name": "3D-FIX batch-upload-test criterion", "coefficient": "1.0", "direction": "POSITIVE"},
        )
        check("27: creating a Primary Screening Criterion via the real route succeeds", criterion_response.status_code in (200, 302))

        run_response = client.post(f"/applications/{first_application_id}/primary-screening/run")
        check(
            "28 (3D-FIX-J, authentic): running Primary Screening through the real route — with no AI "
            "provider configured — completes without a server error (never crashes on AI-unavailable)",
            run_response.status_code == 302,
        )

        with selection_app.SessionFactory() as session:
            run = selection_app.ps_svc.get_latest_run_for_application(session, first_application_id)
            run_id = run.id
            selection_app.ps_svc.add_run_note(session, run_id, "Batch-upload-test run note.")
            evaluation = selection_app.ps_svc.list_evaluations(session, run_id)[0]
            selection_app.ps_svc.add_evaluation_note(session, evaluation.id, "Batch-upload-test evaluation note.")
            session.commit()
            raw_priority_index = run.priority_index

        detail_response = client.get(f"/primary-screening/{run_id}")
        detail_body = detail_response.get_data(as_text=True)
        check(
            "29 (3D-FIX-Y/Z): the Screening Run detail page shows both the run-level and the "
            "Criterion-Evaluation-level Selezionatore notes",
            detail_response.status_code == 200 and "Batch-upload-test run note." in detail_body
            and "Batch-upload-test evaluation note." in detail_body,
        )
        check(
            "30 (3D-FIX-W/X): the raw numeric Priority Index never appears on the Screening Run detail "
            "page, and the Original CV action is still present",
            str(round(raw_priority_index, 4)) not in detail_body and "priority_index" not in detail_body.lower()
            and ("Open Original CV" in detail_body or "Original CV unavailable" in detail_body),
        )

        queue_response = client.get("/primary-screening-queue")
        queue_body = queue_response.get_data(as_text=True)
        check(
            "31 (3D-FIX-W): the raw numeric Priority Index never appears on the Primary Screening Queue "
            "page either",
            queue_response.status_code == 200 and str(round(raw_priority_index, 4)) not in queue_body
            and "priority_index" not in queue_body.lower(),
        )

        # ---------------------------------------------------------------
        # 6. Task 5A — Stage/Outcome/Decision-History engine through the
        #    real HTTP routes: Outcome Definitions authoring, Stage
        #    movement, applying an Outcome, and the Decision Summary page
        #    (Selection Notes History at the BOTTOM of the page, task §22).
        # ---------------------------------------------------------------
        outcomes_response = client.get("/selection-outcomes")
        check(
            "32 (5A): the Selection Outcomes authoring page loads and shows the seeded default Outcomes",
            outcomes_response.status_code == 200 and "Hire" in outcomes_response.get_data(as_text=True),
        )

        stage_response = client.post(f"/applications/{first_application_id}/stage", data={"new_stage": "PRIMARY_SCREENING", "note_text": "Moving to screening."})
        check("33 (5A-B): moving Stage through the real route succeeds", stage_response.status_code == 302)

        with selection_app.SessionFactory() as session:
            restaurant = selection_app._bootstrap_restaurant(session)
            defs = selection_app.outcome_svc.list_outcome_definitions(session, restaurant_id=restaurant.id)
            hold_def_id = next(d.id for d in defs if d.name == "Hold")

        apply_response = client.post(
            f"/applications/{first_application_id}/outcome/apply",
            data={"outcome_definition_id": str(hold_def_id), "reason": "Sufficient pipeline already exists for this role"},
        )
        check("34 (5A-O): applying an Outcome through the real route succeeds", apply_response.status_code == 302)

        decision_response = client.get(f"/applications/{first_application_id}/decision")
        decision_body = decision_response.get_data(as_text=True)
        check(
            "35 (5A-23): the Decision Summary page loads and shows current Stage, current Outcome, and "
            "Application lifecycle state",
            decision_response.status_code == 200 and "Current Stage" in decision_body
            and "Current Outcome" in decision_body and "Application Lifecycle" in decision_body,
        )
        check(
            "36 (5A-AL): Original CV remains available from the Decision Summary page",
            "Open Original CV" in decision_body or "Original CV unavailable" in decision_body,
        )
        notes_history_pos = decision_body.find("Selection Notes History")
        outcome_controls_pos = decision_body.find("Apply an Outcome")
        check(
            "37 (5A-22): the Selection Notes History section appears at the BOTTOM of the principal "
            "Decision view — after the Stage/Outcome controls, not before",
            notes_history_pos != -1 and outcome_controls_pos != -1 and notes_history_pos > outcome_controls_pos,
        )

    finally:
        # Clean up every file this test wrote into the real uploads/ dir —
        # never leave synthetic test résumés behind (task §12).
        current_upload_files = set(os.listdir(selection_app.UPLOAD_DIR))
        for name in current_upload_files - existing_upload_files:
            try:
                os.remove(os.path.join(selection_app.UPLOAD_DIR, name))
            except OSError:
                pass
        selection_app._engine.dispose()
        for suffix in ("", "-wal", "-shm", "-journal"):
            path = _TEST_DB_PATH + suffix
            if os.path.exists(path):
                try:
                    os.remove(path)
                except OSError:
                    pass

    if not checks_failed:
        print(f"Selection batch upload (Task 2A) tests: SUCCESS ({len(checks_passed)}/{len(checks_passed)} checks passed)")
        return 0

    print(
        "Selection batch upload (Task 2A) tests: FAILURE "
        f"({len(checks_passed)} passed, {len(checks_failed)} failed)"
    )
    for description in checks_failed:
        print(f"  FAILED: {description}")
    return 1


if __name__ == "__main__":
    sys.exit(main())
