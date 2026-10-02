#!/usr/bin/env python
"""HTTP test for Reconciliation Standards and the inline exception editor
(BANK_RECONCILIATION_STANDARDS_001).

The Product Owner's 50 points (transaction WHAT, Standard schema, compact
row, Automatic, expansion, exceptions, new Standards, conflicts, status,
Export, regression) plus A-L on accounting-destination provenance and
Reclassify. "Future" transactions arrive through the real CSV upload, so
Automatic is proven through the actual import path.

Runs against DISPOSABLE SQLite databases created here and deleted at the end.
"""

from __future__ import annotations

import io
import json
import os
import re
import sqlite3
import sys
import tempfile
from datetime import date

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
APP_DIR = os.path.normpath(os.path.join(BASE_DIR, ".."))
if APP_DIR not in sys.path:
    sys.path.insert(0, APP_DIR)

_TEST_DB_FD, _TEST_DB_PATH = tempfile.mkstemp(suffix=".db", prefix="rfoneweb_bank_std_")
os.close(_TEST_DB_FD)
os.remove(_TEST_DB_PATH)
os.environ["RFONE_DATABASE_URL"] = f"sqlite:///{_TEST_DB_PATH.replace(os.sep, '/')}"
os.environ["RFONE_WEB_TEST_SECRET_KEY"] = "bank-std-test-secret"

_DATA_STORE_DIR = os.path.normpath(os.path.join(APP_DIR, "..", "RF-One Data Store"))
if _DATA_STORE_DIR not in sys.path:
    sys.path.insert(0, _DATA_STORE_DIR)

from rfone_data_store.database import run_migrations_to_head  # noqa: E402
run_migrations_to_head(os.environ["RFONE_DATABASE_URL"])

import app as web_app  # noqa: E402
from db import SessionFactory  # noqa: E402
from sqlalchemy import func, select  # noqa: E402
from sqlalchemy.exc import IntegrityError  # noqa: E402
from werkzeug.datastructures import MultiDict  # noqa: E402
from rfone_data_store import models as m  # noqa: E402
from rfone_data_store import rfone_account_service as account_service  # noqa: E402
from rfone_data_store.bank_reconciliation import card_configuration  # noqa: E402
from rfone_data_store.bank_reconciliation import configuration as config_service  # noqa: E402
from rfone_data_store.bank_reconciliation import export as export_service  # noqa: E402
from rfone_data_store.bank_reconciliation import monthly_source  # noqa: E402
from rfone_data_store.bank_reconciliation import receiver_candidates  # noqa: E402
from rfone_data_store.bank_reconciliation import recognition  # noqa: E402
from rfone_data_store.bank_reconciliation import reconciliation_standards  # noqa: E402

CSRF_RE = re.compile(r'name="csrf_token" value="([^"]+)"')
BANK_HEADER = "Details,Posting Date,Description,Amount,Type,Balance,Check or Slip #\n"
CARD_HEADER = "Card,Transaction Date,Post Date,Description,Category,Type,Amount,Memo\n"


def bank_csv(rows):
    return (BANK_HEADER + "".join(f"DEBIT,{d},{t},{a},ACH_DEBIT,1000.00,\n" for d, t, a in rows)).encode()


def card_csv(rows):
    return (CARD_HEADER + "".join(f"3333,{d},{d},{t},Food,Sale,{a},\n" for d, t, a in rows)).encode()


def main() -> int:
    passed: list[str] = []
    failed: list[str] = []
    extra: list[str] = []

    def check(description: str, condition: bool, detail: str = "") -> None:
        if condition:
            passed.append(description)
            print(f"  PASS  {description}")
        else:
            failed.append(description)
            print(f"  FAIL  {description}" + (f"   [{str(detail)[:400]}]" if detail else ""))

    try:
        # ------------------------------------------------------------ fixture
        ids: dict[str, int] = {}
        with SessionFactory() as db:
            account_service.create_account(db, username="st_operator", display_name="Pino Miraglia",
                                           password="OperatorPass123!", status="ACTIVE", is_admin=False)
            db.commit()
            op = db.query(m.RFOneAccount).filter_by(username="st_operator").one()
            account_service.set_domain_access(db, account_id=op.id, domain_code="BANK", enabled=True, role_code=None)
            ids["op"] = op.id
            alpha = config_service.create_entity(db, name="Winter Park", legal_name="RF Winter Park, LLC")
            dora = config_service.create_entity(db, name="Mount Dora", legal_name="RF Mount Dora, LLC")
            gelati = config_service.create_entity(db, name="Gelati", legal_name="RF Gelati, LLC")
            checking = m.PaymentInstrument(display_name="Chase Operating", instrument_type="BANK_ACCOUNT",
                                           institution="CHASE", last_four="0214",
                                           legal_entity_id=alpha.legal_entity_id, status="ACTIVE")
            card = m.PaymentInstrument(display_name="Chase Ink", instrument_type="CREDIT_CARD",
                                       institution="CHASE", last_four="3333", status="ACTIVE")
            db.add_all([checking, card])
            db.flush()
            card_configuration.assign_settlement_account(
                db, credit_card_payment_instrument_id=card.id, settlement_bank_account_id=checking.id,
                valid_from=date(2026, 1, 1))
            monthly_source.set_control_start(db, year=2026, month=1, account_id=op.id)
            pl = db.scalars(select(m.BankTransactionReason).join(
                m.BankAccountingClassification,
                m.BankAccountingClassification.id == m.BankTransactionReason.accounting_classification_id,
            ).where(m.BankAccountingClassification.statement_type == "PROFIT_LOSS",
                    m.BankTransactionReason.status == "ACTIVE").order_by(m.BankTransactionReason.id)).all()
            # four P&L WHY with four DIFFERENT WHAT
            chosen, seen = [], set()
            for r in pl:
                if r.accounting_classification_id not in seen:
                    chosen.append(r)
                    seen.add(r.accounting_classification_id)
                if len(chosen) == 4:
                    break
            food, cleaning, office, spare = chosen
            bs_why = db.scalars(select(m.BankTransactionReason).join(
                m.BankAccountingClassification,
                m.BankAccountingClassification.id == m.BankTransactionReason.accounting_classification_id,
            ).where(m.BankAccountingClassification.statement_type == "BALANCE_SHEET")).first()
            gfs = config_service.save_who(db, occurrence_id=None, name="Gordon Food Service", active=True,
                                          reason_ids=[food.id, cleaning.id], default_reason_id=food.id,
                                          reporting_entity_ids=[alpha.id],
                                          rules=[config_service.RuleInput(None, "GORDON FOOD SERVICE", "PREFIX", True)])
            cintas = config_service.save_who(db, occurrence_id=None, name="Cintas", active=True,
                                             reason_ids=[cleaning.id], default_reason_id=cleaning.id,
                                             reporting_entity_ids=[], rules=[])
            db.commit()
            ids.update(alpha=alpha.id, dora=dora.id, gelati=gelati.id, checking=checking.id, card=card.id,
                       food=food.id, cleaning=cleaning.id, office=office.id, spare=spare.id, bs_why=bs_why.id,
                       food_what=food.accounting_classification_id,
                       cleaning_what=cleaning.accounting_classification_id,
                       office_what=office.accounting_classification_id,
                       spare_what=spare.accounting_classification_id,
                       gfs=gfs.id, cintas=cintas.id)
            rules_before = db.scalar(select(func.count(m.BankRecognitionRule.id)))

        client = web_app.app.test_client()
        client.post("/login", data={"username": "st_operator", "password": "OperatorPass123!",
                                    "csrf_token": CSRF_RE.search(client.get("/login").data.decode()).group(1)})
        token = CSRF_RE.search(client.get("/bank/reconciliation?year=2026&month=8").data.decode()).group(1)

        def upload(content: bytes, name: str, instrument: int):
            r = client.post("/bank/upload", data={
                "files": (io.BytesIO(content), name), "payment_instrument_id": str(instrument), "csrf_token": token,
            }, content_type="multipart/form-data")
            assert r.status_code == 302, r.status_code

        def post(path, data):
            pairs = list(data.items()) if isinstance(data, dict) else list(data)
            pairs.append(("csrf_token", token))
            return client.post(path, data=MultiDict(pairs))

        def page(month=8) -> str:
            return client.get(f"/bank/reconciliation?year=2026&month={month}").data.decode("utf-8")

        def txn_id(text: str, instrument: int | None = None) -> int:
            with SessionFactory() as db:
                q = select(m.FinancialTransaction.id).where(m.FinancialTransaction.description_original == text)
                if instrument is not None:
                    q = q.where(m.FinancialTransaction.payment_instrument_id == instrument)
                return db.scalar(q.order_by(m.FinancialTransaction.id.desc()))

        def row(html, tid) -> str:
            start = html.index(f'<tr id="t-{tid}"')
            return html[start:html.index("</tr>", start)]

        def attr(r, name):
            match = re.search(rf' {name}="([^"]*)"', r)
            return match.group(1) if match else None

        def state(tid):
            with SessionFactory() as db:
                t = db.get(m.FinancialTransaction, tid)
                e = db.get(m.BankTransactionExplanation, t.explanation_id) if t.explanation_id else None
                a = db.scalars(select(m.BankTransactionAllocation).where(
                    m.BankTransactionAllocation.financial_transaction_id == tid)).all()
                return e, list(a)

        def catalog(html):
            return json.loads(re.search(r'<script type="application/json" id="rp-data">(.*?)</script>', html, re.S).group(1))

        def message(resp) -> str:
            html = client.get(resp.headers["Location"]).data.decode("utf-8")
            found = re.findall(r'<div class="(error|info)">(.*?)</div>', html, re.S)
            return " | ".join(f"{k}:{v}" for k, v in found)

        # ================================================ A — migration of existing rows
        fd, old_path = tempfile.mkstemp(suffix=".db", prefix="rfoneweb_bank_std_old_")
        os.close(fd)
        os.remove(old_path)
        extra.append(old_path)
        old_url = f"sqlite:///{old_path.replace(os.sep, '/')}"
        from alembic import command
        from alembic.config import Config
        cfg = Config(os.path.join(_DATA_STORE_DIR, "alembic.ini"))
        cfg.set_main_option("script_location", os.path.join(_DATA_STORE_DIR, "migrations"))
        os.environ["ALEMBIC_DATABASE_URL_OVERRIDE"] = old_url
        try:
            command.upgrade(cfg, "f6c2e8a4b1d7")
            conn = sqlite3.connect(old_path)

            def insert(table, values):
                info = conn.execute(f"PRAGMA table_info({table})").fetchall()
                row_values = dict(values)
                for _cid, name, ctype, notnull, default, pk in info:
                    if name in row_values or not notnull or default is not None or pk:
                        continue
                    row_values[name] = 1 if "INT" in ctype.upper() or "BOOL" in ctype.upper() else (
                        "2026-08-01" if "DATE" in ctype.upper() else "X")
                cols = ", ".join(row_values)
                conn.execute(f"INSERT INTO {table} ({cols}) VALUES ({', '.join('?' * len(row_values))})",
                             list(row_values.values()))
            insert("payment_instruments", {"id": 1, "display_name": "Old", "instrument_type": "BANK_ACCOUNT", "status": "ACTIVE"})
            insert("financial_transactions", {"id": 1, "payment_instrument_id": 1, "amount_minor": -100,
                                              "status": "COMPLETED", "classification": "UNKNOWN"})
            insert("bank_transaction_explanations", {"id": 1, "financial_transaction_id": 1, "decision_status": "SUGGESTED"})
            insert("bank_transaction_allocations", {"id": 1, "financial_transaction_id": 1, "allocation_index": 0,
                                                    "amount_minor": -100, "status": "UNALLOCATED"})
            conn.commit()
            conn.close()
            command.upgrade(cfg, "head")
        finally:
            os.environ.pop("ALEMBIC_DATABASE_URL_OVERRIDE", None)
        conn = sqlite3.connect(old_path)
        migrated = (conn.execute("SELECT accounting_destination_source, reconciliation_standard_id FROM bank_transaction_explanations").fetchall(),
                    conn.execute("SELECT accounting_destination_source, reconciliation_standard_id FROM bank_transaction_allocations").fetchall())
        head = conn.execute("SELECT version_num FROM alembic_version").fetchone()[0]
        conn.close()
        check("A. existing decision and allocation rows migrate to source WHY with no Standard lineage",
              migrated == ([("WHY", None)], [("WHY", None)]) and head == "a7c3e9d5f2b8", f"{migrated} {head}")

        # ================================================ first month of real imports
        upload(bank_csv([
            ("08/05/2026", "GORDON FOOD SERVICE #1234 ORLANDO", "-1284.37"),
            ("08/06/2026", "GORDON FOOD SERVICE #5678 ORLANDO", "-212.10"),
            ("08/07/2026", "LINEN CO WEEKLY 7788", "-50.00"),
            ("08/08/2026", "ZELLE PAYMENT TO MARIA LOPEZ", "-300.00"),
            ("08/09/2026", "TIPS PAYABLE SETTLEMENT", "-80.00"),
        ]), "chase_0214_aug.csv", ids["checking"])
        t_gfs1 = txn_id("GORDON FOOD SERVICE #1234 ORLANDO")
        t_gfs2 = txn_id("GORDON FOOD SERVICE #5678 ORLANDO")
        t_linen = txn_id("LINEN CO WEEKLY 7788")
        t_zelle = txn_id("ZELLE PAYMENT TO MARIA LOPEZ")
        t_bs = txn_id("TIPS PAYABLE SETTLEMENT")
        html = page()
        r_gfs1, r_zelle = row(html, t_gfs1), row(html, t_zelle)
        check("39. Needs review shows a red X (with its text)",
              'class="rp-st rp-st-review" role="img" aria-label="Needs review"' in r_zelle and "Needs review</span>" in r_zelle)
        check("13 (prev). ordinary WHO recognition alone does not create Automatic",
              attr(r_gfs1, "data-who") == str(ids["gfs"]) and attr(r_gfs1, "data-status") == "Needs review")

        # ================================================ 19-24 expansion
        rows_count = html.count('data-row-id="')
        compact = html[html.index('<tbody id="rp-body">'):html.index("</tbody>", html.index('<tbody id="rp-body">'))]
        check("19. every row has an inline Expand control", compact.count('data-action="expand"') == rows_count == 5)
        data = catalog(html)
        with SessionFactory() as db:
            active_whos = db.scalar(select(func.count()).where(m.BankOccurrence.status == "ACTIVE"))
            active_whys = db.scalar(select(func.count()).where(m.BankTransactionReason.status == "ACTIVE"))
            pl_whats = sum(1 for c in db.scalars(select(m.BankAccountingClassification)) if c.is_what)
            active_entities = db.scalar(select(func.count()).where(m.ReportingEntity.status == "ACTIVE"))
        check("20. the editor's WHO combo has every active WHO", len(data["whos"]) == active_whos)
        check("21. the editor's WHY combo has every active WHY (not only the WHO's)", len(data["whys"]) == active_whys)
        check("22. the editor's WHAT combo has every P&L WHAT", len(data["whats"]) == pl_whats)
        check("23. the editor's For Whom combo has every active entity",
              sum(1 for e in data["entities"] if e["active"]) == active_entities)
        check("24. current values travel with the row for preselection (WHO, WHY, WHAT, For Whom)",
              attr(r_gfs1, "data-who") == str(ids["gfs"]) and 'data-saved="' in r_gfs1
              and html.count('id="rp-editor-row"') == 1)

        # ================================================ B + transaction WHAT
        resp = post(f"/bank/reconciliation/{t_gfs1}/save", {
            "mode": "keep", "occurrence_id": ids["gfs"], "why_id": ids["food"], "what_id": ids["food_what"],
            "for_whom_id": ids["alpha"]})
        e1, a1 = state(t_gfs1)
        check("B. a human choice equal to the WHY's destination persists as source WHY (decision and allocation)",
              e1.accounting_destination_source == "WHY" and a1[0].accounting_destination_source == "WHY"
              and e1.accounting_classification_id == ids["food_what"] and a1[0].reconciliation_standard_id is None)
        check("41. Confirmed shows a green check",
              'rp-st rp-st-confirmed' in row(page(), t_gfs1) and 'data-status="Confirmed"' in row(page(), t_gfs1))
        resp = post(f"/bank/reconciliation/{t_gfs2}/save", {
            "mode": "keep", "occurrence_id": ids["gfs"], "why_id": ids["food"], "what_id": ids["office_what"],
            "for_whom_id": ids["alpha"]})
        gfs2_msg = message(resp)
        e2, a2 = state(t_gfs2)
        with SessionFactory() as db:
            master = db.get(m.BankTransactionReason, ids["food"]).accounting_classification_id
        check("1 / C / 4. an explicit P&L WHAT different from the WHY's persists as a TRANSACTION exception",
              e2.accounting_classification_id == ids["office_what"] and e2.accounting_destination_source == "TRANSACTION"
              and a2[0].accounting_classification_id == ids["office_what"] and a2[0].accounting_destination_source == "TRANSACTION",
              f"{gfs2_msg} {e2 and (e2.accounting_classification_id, e2.accounting_destination_source)} "
              f"{[(x.accounting_classification_id, x.accounting_destination_source, x.status) for x in a2]} {ids['office_what']} {ids['food_what']}")
        check("2 / F. the WHY's master WHAT is unchanged by a transaction exception", master == ids["food_what"])

        # 3 — a preserved Balance Sheet destination
        with SessionFactory() as db:
            bs_dest = db.get(m.BankTransactionReason, ids["bs_why"]).accounting_classification_id
        post(f"/bank/reconciliation/{t_bs}/save", {"mode": "keep", "occurrence_id": ids["gfs"], "why_id": ids["bs_why"],
                                                   "what_id": bs_dest, "for_whom_id": ids["alpha"]})
        e_bs, a_bs = state(t_bs)
        bs_label = catalog(page())["destinations"].get(str(bs_dest), "")
        check("3. an existing Balance Sheet destination is preserved and shown truthfully (not converted)",
              e_bs.accounting_classification_id == bs_dest and e_bs.accounting_destination_source == "WHY"
              and bs_label.startswith("Balance Sheet destination"), bs_label)

        # ================================================ G / H — Reclassify
        with SessionFactory() as db:
            before_rows = db.scalar(select(func.count()).where(m.BankTransactionExplanation.financial_transaction_id == t_gfs1))
            recognition.reclassify_transaction(db, transaction_id=t_gfs1, confirmed_by_account_id=ids["op"])
            db.commit()
            after_rows = db.scalar(select(func.count()).where(m.BankTransactionExplanation.financial_transaction_id == t_gfs1))
            try:
                recognition.reclassify_transaction(db, transaction_id=t_gfs2, confirmed_by_account_id=ids["op"])
                h_refused = False
            except ValueError as exc:
                db.rollback()
                h_refused = "chosen explicitly for this transaction" in str(exc)
            kept = db.get(m.BankTransactionExplanation, db.get(m.FinancialTransaction, t_gfs2).explanation_id)
        check("G. Reclassify re-derives a WHY-source destination", after_rows == before_rows + 1)
        check("H. Reclassify refuses a TRANSACTION destination and leaves it unchanged",
              h_refused and kept.accounting_classification_id == ids["office_what"])

        # ================================================ 12-14 compact row
        post(f"/bank/reconciliation/{t_gfs1}/save", {"mode": "keep", "occurrence_id": ids["gfs"], "why_id": ids["food"],
                                                     "what_id": ids["food_what"], "for_whom_id": ids["alpha"]})
        post(f"/bank/reconciliation/{t_zelle}/who", {"occurrence_id": ids["cintas"], "why_id": ids["cleaning"]})
        resp = post(f"/bank/reconciliation/{t_zelle}/confirm", {"for_whom_id": ids["dora"]})
        ez, az = state(t_zelle)
        check("12. compact Confirm (row collapsed) confirms with WHO, WHY, WHAT and the chosen For Whom",
              az and az[0].decision_source == "HUMAN" and az[0].reporting_entity_id == ids["dora"]
              and ez.decision_source == "HUMAN" and ez.accounting_classification_id == ids["cleaning_what"]
              and 'data-status="Confirmed"' in row(page(), t_zelle))
        with SessionFactory() as db:
            standards_before = db.scalar(select(func.count(m.BankReconciliationStandard.id)))
        post(f"/bank/reconciliation/{t_linen}/who", {"occurrence_id": ids["cintas"], "why_id": ids["cleaning"]})
        resp = post(f"/bank/reconciliation/{t_linen}/standard", {"for_whom_id": ids["dora"]})
        el, al = state(t_linen)
        with SessionFactory() as db:
            std = db.scalar(select(m.BankReconciliationStandard).order_by(m.BankReconciliationStandard.id.desc()))
            standards_after = db.scalar(select(func.count(m.BankReconciliationStandard.id)))
        check("13. compact Set as Standard (row collapsed) creates a Standard", standards_after == standards_before + 1, message(resp))
        check("14 / J / 3 (prev). the row that created the Standard is human CONFIRMED, not Automatic, with source WHY",
              al[0].decision_source == "HUMAN" and al[0].reconciliation_standard_id is None
              and el.accounting_destination_source == "WHY" and el.reconciliation_standard_id is None
              and 'data-status="Confirmed"' in row(page(), t_linen))
        check("4 (prev). Standard signature: no recognition rule -> EXACT normalized description, this account, this direction",
              std.match_type == "EXACT_NORMALIZED_DESCRIPTION" and std.normalized_pattern == "LINEN CO WEEKLY 7788"
              and std.match_field == "DESCRIPTION" and std.payment_instrument_id == ids["checking"] and std.direction == "DEBIT")
        check("5. Standard persists WHO", std.occurrence_id == ids["cintas"])
        check("6. Standard persists WHY", std.transaction_reason_id == ids["cleaning"])
        check("7. Standard persists WHAT", std.accounting_classification_id == ids["cleaning_what"])
        check("8. Standard persists For Whom", std.reporting_entity_id == ids["dora"])
        check("9. Standard audit persists (who approved, when, from which transaction, ACTIVE)",
              std.approved_by_account_id == ids["op"] and std.approved_at is not None
              and std.created_from_transaction_id == t_linen and std.status == "ACTIVE")
        ids["linen_std"] = std.id

        # Case A: the GFS rows carry a recognition rule -> its PREFIX is reused.
        resp = post(f"/bank/reconciliation/{t_gfs2}/save", {
            "mode": "standard", "occurrence_id": ids["gfs"], "why_id": ids["food"], "what_id": ids["office_what"],
            "for_whom_id": ids["alpha"]})
        with SessionFactory() as db:
            gfs_std = db.scalar(select(m.BankReconciliationStandard).where(
                m.BankReconciliationStandard.normalized_pattern == "GORDON FOOD SERVICE"))
        check("Signature case A: the recognition rule's PREFIX pattern is reused, scoped to account and direction",
              gfs_std is not None and gfs_std.match_type == "PREFIX" and gfs_std.payment_instrument_id == ids["checking"]
              and gfs_std.direction == "DEBIT" and gfs_std.accounting_classification_id == ids["office_what"], message(resp))
        ids["gfs_std"] = gfs_std.id

        # ================================================ 15-18 / D / K Automatic on the next import
        upload(bank_csv([
            ("08/20/2026", "GORDON FOOD SERVICE #9999 ORLANDO", "-400.00"),
            ("08/21/2026", "LINEN CO WEEKLY 7788", "-55.00"),
            ("08/22/2026", "LINEN CO WEEKLY 7788", "-60.00"),
            ("08/23/2026", "LINEN CO WEEKLY 7788", "-65.00"),
        ]), "chase_0214_aug_b.csv", ids["checking"])
        upload(card_csv([("08/24/2026", "LINEN CO WEEKLY 7788", "-70.00"),
                         ("08/25/2026", "GORDON FOOD SERVICE #4444 ORLANDO", "-90.00")]), "chase_3333_aug.csv", ids["card"])
        with SessionFactory() as db:
            new_ids = dict(db.execute(select(m.FinancialTransaction.description_original, func.max(m.FinancialTransaction.id))
                                      .where(m.FinancialTransaction.posting_date >= date(2026, 8, 20))
                                      .group_by(m.FinancialTransaction.description_original, m.FinancialTransaction.posting_date)).all())
            linen_auto = db.scalars(select(m.FinancialTransaction.id).where(
                m.FinancialTransaction.description_original == "LINEN CO WEEKLY 7788",
                m.FinancialTransaction.payment_instrument_id == ids["checking"],
                m.FinancialTransaction.posting_date >= date(2026, 8, 21)).order_by(m.FinancialTransaction.id)).all()
        t_gfs_auto = txn_id("GORDON FOOD SERVICE #9999 ORLANDO")
        t_card_linen = txn_id("LINEN CO WEEKLY 7788", ids["card"])
        t_card_gfs = txn_id("GORDON FOOD SERVICE #4444 ORLANDO", ids["card"])
        ea, aa = state(t_gfs_auto)
        html = page()
        r_auto = row(html, t_gfs_auto)
        check("15. a later matching transaction becomes Automatic without a click",
              attr(r_auto, "data-status") == "Automatic")
        check("16 / 40. Automatic shows an orange check (with its text)",
              'rp-st rp-st-automatic' in r_auto and "Automatic</span>" in r_auto)
        check("17. the complete WHO / WHY / WHAT / For Whom of the Standard is applied",
              ea.occurrence_id == ids["gfs"] and ea.transaction_reason_id == ids["food"]
              and ea.accounting_classification_id == ids["office_what"] and aa[0].reporting_entity_id == ids["alpha"]
              and aa[0].status == "COMPLETE" and aa[0].decision_source == "RULE")
        check("D / 10 / 11. Automatic destination is source STANDARD with lineage on BOTH decision and allocation",
              ea.accounting_destination_source == "STANDARD" and aa[0].accounting_destination_source == "STANDARD"
              and ea.reconciliation_standard_id == ids["gfs_std"] and aa[0].reconciliation_standard_id == ids["gfs_std"])
        check("K. the future Standard-applied row is Automatic, not Confirmed", attr(r_auto, "data-status") == "Automatic")
        check("scope: the same text on another card is not matched by an account-scoped Standard",
              attr(row(html, t_card_linen), "data-status") == "Needs review")
        std_info = catalog(html)["standards"].get(str(ids["gfs_std"]), {})
        check("23 (prev). an Automatic row identifies its Standard for the expanded detail (readable, no ids)",
              attr(r_auto, "data-standard") == str(ids["gfs_std"]) and std_info.get("pattern") == "GORDON FOOD SERVICE"
              and std_info.get("scope") == "Chase Operating ··0214" and std_info.get("approved_by") == "Pino Miraglia"
              and std_info.get("for_whom") == "Winter Park", str(std_info))

        # ================================================ I / L
        with SessionFactory() as db:
            try:
                recognition.reclassify_transaction(db, transaction_id=t_gfs_auto, confirmed_by_account_id=ids["op"])
                i_refused = False
            except ValueError as exc:
                db.rollback()
                i_refused = "Standard" in str(exc)
        check("I. Reclassify refuses a STANDARD destination", i_refused)
        with SessionFactory() as db:
            config_service.update_why(db, transaction_reason_id=ids["cleaning"],
                                      name=db.get(m.BankTransactionReason, ids["cleaning"]).name, description=None,
                                      what_id=ids["spare_what"], active=True)
            db.commit()
        upload(bank_csv([("08/26/2026", "LINEN CO WEEKLY 7788", "-75.00")]), "chase_0214_aug_c.csv", ids["checking"])
        t_linen_late = txn_id("LINEN CO WEEKLY 7788", ids["checking"])
        el2, al2 = state(t_linen_late)
        check("L. a Standard's destination stays stable when the WHY's master mapping changes later",
              el2.accounting_classification_id == ids["cleaning_what"] and al2[0].accounting_classification_id == ids["cleaning_what"]
              and el2.reconciliation_standard_id == ids["linen_std"])

        # ================================================ 25-31 exceptions
        with SessionFactory() as db:
            gfs_std_before = db.get(m.BankReconciliationStandard, ids["gfs_std"])
            snapshot_before = (gfs_std_before.occurrence_id, gfs_std_before.transaction_reason_id,
                               gfs_std_before.accounting_classification_id, gfs_std_before.reporting_entity_id,
                               gfs_std_before.status, gfs_std_before.normalized_pattern)
        resp = post(f"/bank/reconciliation/{t_gfs_auto}/save", {
            "mode": "keep", "occurrence_id": ids["gfs"], "why_id": ids["food"], "what_id": ids["office_what"],
            "for_whom_id": ids["dora"]})
        ex, ax = state(t_gfs_auto)
        with SessionFactory() as db:
            s = db.get(m.BankReconciliationStandard, ids["gfs_std"])
            snapshot_after = (s.occurrence_id, s.transaction_reason_id, s.accounting_classification_id,
                              s.reporting_entity_id, s.status, s.normalized_pattern)
        check("25 / 26. an Automatic row can change only For Whom; Save keeps the change to this transaction",
              ax[0].reporting_entity_id == ids["dora"] and ax[0].decision_source == "HUMAN"
              and ax[0].reconciliation_standard_id is None and ex.reconciliation_standard_id is None
              and ex.accounting_destination_source == "TRANSACTION")
        check("27. the Standard is unchanged", snapshot_before == snapshot_after)
        check("28. the result is Confirmed", 'data-status="Confirmed"' in row(page(), t_gfs_auto))
        post(f"/bank/reconciliation/{linen_auto[0]}/save", {
            "mode": "keep", "occurrence_id": ids["gfs"], "why_id": ids["office"], "what_id": ids["food_what"],
            "for_whom_id": ids["gelati"]})
        ew, aw = state(linen_auto[0])
        check("29 / 30 / 31. an Automatic row can change WHO, WHY (any WHY) and WHAT",
              ew.occurrence_id == ids["gfs"] and ew.transaction_reason_id == ids["office"]
              and ew.accounting_classification_id == ids["food_what"] and aw[0].reporting_entity_id == ids["gelati"]
              and 'data-status="Confirmed"' in row(page(), linen_auto[0]))

        # ================================================ 36-38 conflicts
        with SessionFactory() as db:
            count_before = db.scalar(select(func.count(m.BankReconciliationStandard.id)))
            decision_before = db.get(m.FinancialTransaction, linen_auto[1]).explanation_id
        resp = post(f"/bank/reconciliation/{linen_auto[1]}/save", {
            "mode": "standard", "occurrence_id": ids["cintas"], "why_id": ids["cleaning"],
            "what_id": ids["cleaning_what"], "for_whom_id": ids["gelati"]})
        conflict_msg = message(resp)
        with SessionFactory() as db:
            count_after = db.scalar(select(func.count(m.BankReconciliationStandard.id)))
            decision_after = db.get(m.FinancialTransaction, linen_auto[1]).explanation_id
        check("36. same signature with a different result is refused, and nothing is written",
              count_after == count_before and decision_after == decision_before
              and "no distinct recognition pattern" in conflict_msg, conflict_msg)
        post(f"/bank/reconciliation/{linen_auto[1]}/save", {
            "mode": "keep", "occurrence_id": ids["cintas"], "why_id": ids["cleaning"],
            "what_id": ids["cleaning_what"], "for_whom_id": ids["gelati"]})
        _e, a37 = state(linen_auto[1])
        check("37. that exception can still be saved individually", a37 and a37[0].reporting_entity_id == ids["gelati"]
              and 'data-status="Confirmed"' in row(page(), linen_auto[1]))
        resp = post(f"/bank/reconciliation/{linen_auto[2]}/save", {
            "mode": "standard", "occurrence_id": ids["cintas"], "why_id": ids["cleaning"],
            "what_id": ids["cleaning_what"], "for_whom_id": ids["dora"]})
        with SessionFactory() as db:
            count_idem = db.scalar(select(func.count(m.BankReconciliationStandard.id)))
        check("38. same signature with the same result reuses the Standard (idempotent)",
              count_idem == count_after and 'data-status="Confirmed"' in row(page(), linen_auto[2]), message(resp))

        # ================================================ 32-35 new distinct Standard
        resp = post(f"/bank/reconciliation/{t_card_gfs}/save", {
            "mode": "standard", "occurrence_id": ids["gfs"], "why_id": ids["cleaning"],
            "what_id": ids["spare_what"], "for_whom_id": ids["gelati"]})
        with SessionFactory() as db:
            card_std = db.scalar(select(m.BankReconciliationStandard).where(
                m.BankReconciliationStandard.payment_instrument_id == ids["card"]))
            s = db.get(m.BankReconciliationStandard, ids["gfs_std"])
            original_after = (s.occurrence_id, s.transaction_reason_id, s.accounting_classification_id,
                              s.reporting_entity_id, s.status, s.normalized_pattern)
        check("32. Save as New Standard leaves the original Standard unchanged", original_after == snapshot_before)
        check("33. a new Standard is created when the signature is distinguishable (another card)",
              card_std is not None and card_std.normalized_pattern == "GORDON FOOD SERVICE" and card_std.match_type == "PREFIX",
              message(resp))
        check("34. the new Standard stores the complete changed result",
              card_std is not None and (card_std.occurrence_id, card_std.transaction_reason_id,
                                        card_std.accounting_classification_id, card_std.reporting_entity_id)
              == (ids["gfs"], ids["cleaning"], ids["spare_what"], ids["gelati"]))
        upload(card_csv([("08/28/2026", "GORDON FOOD SERVICE #4545 ORLANDO", "-95.00")]), "chase_3333_aug_b.csv", ids["card"])
        t_card_future = txn_id("GORDON FOOD SERVICE #4545 ORLANDO", ids["card"])
        ecf, acf = state(t_card_future)
        check("35. a future variant becomes Automatic through the new Standard",
              acf and acf[0].reconciliation_standard_id == card_std.id and 'data-status="Automatic"' in row(page(), t_card_future))

        # ambiguous / conflicting Standards -> Needs review
        with SessionFactory() as db:
            reconciliation_standards.create_or_reuse(
                db, signature=reconciliation_standards.Signature("CONTAINS_TEXT", "WEEKLY", "DESCRIPTION", ids["checking"], "DEBIT"),
                result=reconciliation_standards.Result(ids["gfs"], ids["food"], ids["food_what"], ids["alpha"]),
                approved_by_account_id=ids["op"], created_from_transaction_id=None)
            db.commit()
        upload(bank_csv([("08/29/2026", "LINEN CO WEEKLY 7788", "-80.00")]), "chase_0214_aug_d.csv", ids["checking"])
        t_conflict = txn_id("LINEN CO WEEKLY 7788", ids["checking"])
        e_c, a_c = state(t_conflict)
        check("14 / 15 (prev). two matching Standards with different results leave the row Needs review (no guess)",
              not a_c and 'data-status="Needs review"' in row(page(), t_conflict))

        # ================================================ E — lineage constraint
        with SessionFactory() as db:
            bad = db.get(m.BankTransactionExplanation, db.get(m.FinancialTransaction, t_gfs1).explanation_id)
            bad.accounting_destination_source = "STANDARD"
            try:
                db.commit()
                e_refused = False
            except IntegrityError:
                db.rollback()
                e_refused = True
            alloc = db.scalars(select(m.BankTransactionAllocation).where(
                m.BankTransactionAllocation.financial_transaction_id == t_gfs1)).first()
            alloc.reconciliation_standard_id = ids["gfs_std"]
            try:
                db.commit()
                e2_refused = False
            except IntegrityError:
                db.rollback()
                e2_refused = True
        check("E. STANDARD without a Standard, and a Standard id on a WHY-source row, are refused by the database",
              e_refused and e2_refused)

        # ================================================ 42-46 Export
        with SessionFactory() as db:
            blockers = export_service.compute_export_blockers(db, year=2026, month=8)
            review_ids = {int(x) for b in blockers for x in re.findall(r"^Needs review: transaction id=(\d+)", b.reason)}
            statuses = {}
            for tid in (t_gfs1, t_gfs2, t_zelle, t_linen, t_card_future, t_linen_late, t_conflict, t_card_linen):
                statuses[tid] = attr(row(page(), tid), "data-status")
        check("42 / 12 (prev). Needs review blocks the Export", {t_conflict, t_card_linen} <= review_ids)
        check("44 / 10 (prev). Automatic is exportable", t_card_future not in review_ids and t_linen_late not in review_ids
              and statuses[t_card_future] == "Automatic")
        check("45 / 11 (prev). Confirmed is exportable", not ({t_gfs1, t_gfs2, t_zelle, t_linen} & review_ids))
        check("43. an unsaved edit is never persisted: the stored row stays what was approved (browser test covers the "
              "on-screen Needs review)", statuses[t_linen_late] == "Automatic")
        # a month where everything is reconciled builds the unchanged workbook
        upload(bank_csv([("09/03/2026", "LINEN CO WEEKLY 7788", "-81.00")]), "chase_0214_sep.csv", ids["checking"])
        with SessionFactory() as db:
            # retire the conflicting CONTAINS Standard so September is unambiguous
            for s in db.scalars(select(m.BankReconciliationStandard).where(m.BankReconciliationStandard.match_type == "CONTAINS_TEXT")):
                s.status = "INACTIVE"
            db.commit()
        upload(bank_csv([("09/04/2026", "LINEN CO WEEKLY 7788", "-82.00")]), "chase_0214_sep_b.csv", ids["checking"])
        t_sep1 = txn_id("LINEN CO WEEKLY 7788", ids["checking"])
        with SessionFactory() as db:
            sep_first = db.scalar(select(m.FinancialTransaction.id).where(
                m.FinancialTransaction.posting_date == date(2026, 9, 3)))
        post(f"/bank/reconciliation/{sep_first}/save", {"mode": "keep", "occurrence_id": ids["cintas"], "why_id": ids["cleaning"],
                                                        "what_id": ids["cleaning_what"], "for_whom_id": ids["dora"]})
        with SessionFactory() as db:
            sep_blockers = export_service.compute_export_blockers(db, year=2026, month=9)
            workbook = export_service.build_kermali_workbook(db, year=2026, month=9)
        from openpyxl import load_workbook
        sheet = load_workbook(io.BytesIO(workbook)).worksheets[0]
        header = tuple(c.value for c in next(sheet.iter_rows(min_row=1, max_row=1)))
        check("46. with every row Confirmed or Automatic the export runs and its format is unchanged",
              not sep_blockers and header[:len(export_service.KERMALI_COLUMNS)] == tuple(export_service.KERMALI_COLUMNS)
              and attr(row(page(9), t_sep1), "data-status") == "Automatic",
              f"{[b.reason for b in sep_blockers]} {header}")

        # ================================================ 47-50 regression + one modal
        with SessionFactory() as db:
            rules_after = db.scalar(select(func.count(m.BankRecognitionRule.id)))
            purpose_rules = db.scalar(select(func.count()).where(m.BankRecognitionRule.determines_purpose.is_(True)))
            db.add(m.BankRecognitionRule(match_type="EXACT_NORMALIZED_DESCRIPTION", normalized_pattern="X",
                                         match_field="DESCRIPTION", determines_purpose=True,
                                         occurrence_id=ids["gfs"], status="ACTIVE"))
            try:
                db.commit()
                invariant_kept = False
            except IntegrityError:
                db.rollback()
                invariant_kept = True
        check("47. Standards create no recognition rule; recognition rules stay WHO-only",
              rules_after == rules_before and purpose_rules == 0)
        check("48. the WHO->WHY invariant CHECK still refuses a description rule that decides purpose", invariant_kept)
        check("49. Configuration still renders", client.get("/bank/configuration").status_code == 200)

        def refuse(*a, **k):
            raise AssertionError("Classification must not run")
        saved = receiver_candidates.build_candidates
        receiver_candidates.build_candidates = refuse
        try:
            html = page()
            ok = post(f"/bank/reconciliation/{t_card_linen}/save", {
                "mode": "keep", "occurrence_id": ids["cintas"], "why_id": ids["cleaning"],
                "what_id": ids["cleaning_what"], "for_whom_id": ids["alpha"]}).status_code == 302
        finally:
            receiver_candidates.build_candidates = saved
        compact = html[html.index('<tbody id="rp-body">'):html.index("</tbody>", html.index('<tbody id="rp-body">'))]
        check("50 / 26 (prev). no Classification call", ok)
        check("25 (prev). only one reusable WHO modal and one inline editor",
              html.count('id="rp-who-modal"') == 1 and html.count('id="rp-editor-row"') == 1)
        check("27 (prev). no WHAT on the compact rows", "WHAT" not in compact)
    finally:
        for path in [_TEST_DB_PATH] + extra:
            try:
                os.remove(path)
            except OSError:
                pass

    print(f"\n{len(passed)} passed, {len(failed)} failed.")
    if failed:
        print("FAILED CHECKS:")
        for description in failed:
            print(f" - {description}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
