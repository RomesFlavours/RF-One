"""Monthly Export per Owning Entity, people included (BANK_EXPORT_BY_OWNER_001).

On a throwaway database, over HTTP, opening every generated file:

* the Owning Entity is chosen next to year and month (companies and people,
  each once) and is required before Check and Download;
* a Company's report and a person's report contain only the month's
  movements of that owner's Sources — a card counting for the owner of the
  account it settled to on that date — each movement once;
* another owner's problem does not block the selected report; a Source owned
  by a person never blocks; Sources with no owner are listed separately;
* a payment from a person's account for MD stays in the person's report with
  MD as For Whom;
* an owner with no movement in the month is not READY;
* Check and the file agree: row count and totals per currency;
* nothing is written to the data by Check or Download.

Never touches a real bank file, AWS, or any production database.
"""

from __future__ import annotations

import io
import os
import re
import sys
import tempfile
from datetime import date

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
APP_DIR = os.path.normpath(os.path.join(BASE_DIR, ".."))
if APP_DIR not in sys.path:
    sys.path.insert(0, APP_DIR)

_TEST_DB_FD, _TEST_DB_PATH = tempfile.mkstemp(suffix=".db", prefix="rfoneweb_export_by_owner_")
os.close(_TEST_DB_FD)
os.remove(_TEST_DB_PATH)
os.environ["RFONE_DATABASE_URL"] = f"sqlite:///{_TEST_DB_PATH.replace(os.sep, '/')}"
os.environ["RFONE_WEB_TEST_SECRET_KEY"] = "bank-export-by-owner-secret"

_DATA_STORE_DIR = os.path.normpath(os.path.join(APP_DIR, "..", "RF-One Data Store"))
if _DATA_STORE_DIR not in sys.path:
    sys.path.insert(0, _DATA_STORE_DIR)

from rfone_data_store.database import run_migrations_to_head  # noqa: E402
run_migrations_to_head(os.environ["RFONE_DATABASE_URL"])

import openpyxl  # noqa: E402
from sqlalchemy import func, select  # noqa: E402

import app as web_app  # noqa: E402
from db import SessionFactory  # noqa: E402
from rfone_data_store import models as m  # noqa: E402
from rfone_data_store import rfone_account_service as account_service  # noqa: E402
from rfone_data_store.bank_reconciliation import card_configuration  # noqa: E402
from rfone_data_store.bank_reconciliation import export as export_service  # noqa: E402
from rfone_data_store.bank_reconciliation import row_reconciliation  # noqa: E402
from rfone_data_store.bank_reconciliation import why_catalog  # noqa: E402

CSRF_RE = re.compile(r'name="csrf_token" value="([^"]+)"')
AUG = "year=2026&month=8"


def main() -> int:
    passed: list[str] = []
    failed: list[str] = []

    def check(description: str, condition: bool, detail: str = "") -> None:
        (passed if condition else failed).append(description)
        print(f"  {'PASS' if condition else 'FAIL'}  {description}" + ("" if condition or not detail else f" ({detail})"))

    with SessionFactory() as s:
        account_service.create_account(s, username="owner_export", display_name="Owner Export",
                                       password="OwnerExport123!", status="ACTIVE", is_admin=False)
        s.flush()
        operator = s.query(m.RFOneAccount).filter_by(username="owner_export").one()
        account_service.set_domain_access(s, account_id=operator.id, domain_code="BANK", enabled=True, role_code=None)
        rfmd_llc = m.LegalEntity(legal_name="RF Mount Dora, LLC", status="ACTIVE")
        angeli_llc = m.LegalEntity(legal_name="Angeli E Demoni, LLC", status="ACTIVE")
        s.add_all([rfmd_llc, angeli_llc])
        s.flush()
        rfmd = m.ReportingEntity(code="RE_RFMD", name="RFMD", entity_type="LEGAL", legal_entity_id=rfmd_llc.id,
                                 status="ACTIVE")
        pino = m.ReportingEntity(code="RE_PINO", name="Pino", entity_type="PERSONAL", status="ACTIVE")
        s.add_all([rfmd, pino])
        s.flush()

        def instrument(name, kind="BANK_ACCOUNT", *, legal=None, person=None, currency="USD", last="0000"):
            i = m.PaymentInstrument(legal_entity_id=legal, owner_reporting_entity_id=person, instrument_type=kind,
                                    display_name=name, institution="CHASE", last_four=last, currency=currency)
            s.add(i)
            s.flush()
            return i

        rfmd_checking = instrument("RFMD Checking", legal=rfmd_llc.id, last="1001")
        angeli_checking = instrument("Angeli Checking", legal=angeli_llc.id, last="1002")
        pino_checking = instrument("Pino Personal", person=pino.id, last="1003")
        pino_eur = instrument("Pino Euro", person=pino.id, currency="EUR", last="1004")
        orphan_checking = instrument("Orphan Account", last="1005")
        rfmd_card = instrument("RFMD Card", "CREDIT_CARD", last="2001")
        moving_card = instrument("Moving Card", "CREDIT_CARD", last="2002")
        lonely_card = instrument("Card Without Settlement", "CREDIT_CARD", last="2003")
        card_configuration.assign_settlement_account(
            s, credit_card_payment_instrument_id=rfmd_card.id, settlement_bank_account_id=rfmd_checking.id,
            valid_from=date(2026, 1, 1))
        card_configuration.assign_settlement_account(
            s, credit_card_payment_instrument_id=moving_card.id, settlement_bank_account_id=angeli_checking.id,
            valid_from=date(2026, 1, 1))
        card_configuration.assign_settlement_account(
            s, credit_card_payment_instrument_id=moving_card.id, settlement_bank_account_id=pino_checking.id,
            valid_from=date(2026, 8, 15))
        lonely_card.linked_instrument_id = None
        counterparty = s.scalars(select(m.BankOccurrenceType).where(m.BankOccurrenceType.code == "COUNTERPARTY")).one()
        vendor = m.BankOccurrence(canonical_name="Test Vendor", occurrence_type_id=counterparty.id, status="ACTIVE")
        s.add(vendor)
        s.flush()
        why = s.scalars(select(m.BankTransactionReason).where(
            m.BankTransactionReason.status == "ACTIVE",
            m.BankTransactionReason.accounting_classification_id.is_not(None)).order_by(m.BankTransactionReason.id)).first()
        why_catalog.associate(s, occurrence_id=vendor.id, transaction_reason_id=why.id, source="HUMAN")
        seq = [0]

        def tx(inst, day, amount, description):
            seq[0] += 1
            t = m.FinancialTransaction(
                payment_instrument_id=inst.id, bank_source="CHASE_BANK_ACCOUNT", posting_date=day,
                transaction_date=day, description_original=description, description_normalized=description,
                amount_minor=amount, status="COMPLETED", duplicate_status="NONE", review_status="REQUIRES_REVIEW",
                accounting_status="CANONICAL", fingerprint=f"owner-{seq[0]}", classification="UNKNOWN")
            s.add(t)
            s.flush()
            return t

        def confirm(t, for_whom):
            row_reconciliation.record_who(s, transaction_id=t.id, occurrence_id=vendor.id, reason_id=why.id,
                                          account_id=operator.id)
            row_reconciliation.confirm_row(s, transaction_id=t.id, reporting_entity_id=for_whom.id,
                                           account_id=operator.id)

        rows = {
            "rfmd_1": tx(rfmd_checking, date(2026, 8, 3), -10000, "RFMD SUPPLIER ONE"),
            "rfmd_2": tx(rfmd_checking, date(2026, 8, 5), -2550, "RFMD SUPPLIER TWO"),
            "rfmd_card": tx(rfmd_card, date(2026, 8, 12), -4999, "RFMD CARD PURCHASE"),
            "pino_for_md": tx(pino_checking, date(2026, 8, 7), -12345, "PINO PAYS FOR MD"),
            "pino_eur": tx(pino_eur, date(2026, 8, 9), -5000, "PINO EURO PAYMENT"),
            "moving_aug10": tx(moving_card, date(2026, 8, 10), -700, "MOVING CARD BEFORE SWITCH"),
            "moving_aug20": tx(moving_card, date(2026, 8, 20), -800, "MOVING CARD AFTER SWITCH"),
            "angeli_open": tx(angeli_checking, date(2026, 8, 11), -900, "ANGELI STILL TO REVIEW"),
            "orphan": tx(orphan_checking, date(2026, 8, 13), -300, "ORPHAN ACCOUNT MOVE"),
            "lonely": tx(lonely_card, date(2026, 8, 14), -400, "CARD WITH NO SETTLEMENT"),
            "rfmd_july": tx(rfmd_checking, date(2026, 7, 30), -111, "RFMD JULY ROW"),
        }
        for key in ("rfmd_1", "rfmd_2", "rfmd_card", "pino_eur", "moving_aug20", "rfmd_july"):
            confirm(rows[key], rfmd if key.startswith("rfmd") else pino)
        confirm(rows["pino_for_md"], rfmd)          # a personal payment made for MD
        s.commit()
        ids = {k: t.id for k, t in rows.items()}
        ids.update(rfmd=rfmd.id, pino=pino.id, angeli_llc=angeli_llc.id, pino_checking=pino_checking.id,
                   orphan_checking=orphan_checking.id)

    def facts():
        with SessionFactory() as s:
            return (s.scalar(select(func.count(m.BankTransactionExplanation.id))),
                    s.scalar(select(func.count(m.BankTransactionAllocation.id))),
                    s.scalar(select(func.count(m.FinancialTransaction.id))),
                    sorted(s.execute(select(m.PaymentInstrument.id, m.PaymentInstrument.legal_entity_id,
                                            m.PaymentInstrument.owner_reporting_entity_id,
                                            m.PaymentInstrument.linked_instrument_id)).all()))

    before = facts()
    client = web_app.app.test_client()
    csrf = CSRF_RE.search(client.get("/login").data.decode()).group(1)
    client.post("/login", data={"username": "owner_export", "password": "OwnerExport123!", "csrf_token": csrf})

    RFMD, PINO, ANGELI = f"re:{ids['rfmd']}", f"re:{ids['pino']}", f"le:{ids['angeli_llc']}"
    page = client.get(f"/bank/export?{AUG}").data.decode()
    options = re.findall(r'<option value="(re:\d+|le:\d+)"[^>]*>([^<]+)</option>', page)
    check("1. Owning Entity sits next to year and month: companies and people, each once",
          sorted(options) == sorted([(ANGELI, "Angeli E Demoni, LLC (Company)"), (PINO, "Pino (Person)"),
                                     (RFMD, "RFMD (Company)")]), detail=str(options))
    check("1b. without an owner there is no Check result and no Download",
          "export-status" not in page and 'id="export-download"' not in page)
    csrf = CSRF_RE.search(client.get("/bank/review").data.decode()).group(1)
    no_owner = client.post("/bank/export", data={"year": "2026", "month": "8", "csrf_token": csrf})
    check("1c. Download without an owner produces no file",
          no_owner.status_code == 200 and "attachment" not in no_owner.headers.get("Content-Disposition", ""))

    def figures(html):
        return dict(re.findall(r'data-figure="([^"]+)">([^<]+)<', html))

    def download(owner):
        return client.post("/bank/export", data={"year": "2026", "month": "8", "owner": owner, "csrf_token": csrf,
                                                  "download_token": "t123"})

    def open_file(response):
        book = openpyxl.load_workbook(io.BytesIO(response.data))
        bank = [list(r) for r in book["Bank"].iter_rows(values_only=True)]
        report = {r[0]: r[1] for r in book["Report"].iter_rows(values_only=True) if r and len(r) > 1}
        return bank[0], [dict(zip(bank[0], r)) for r in bank[1:]], report

    def matches(html, rows_):
        figs = figures(html)
        totals = {}
        for r in rows_:
            totals.setdefault("USD" if "Euro" not in (r["Account"] or "") else "EUR", 0)
        sums = {}
        for r in rows_:
            cur = "EUR" if r["Account"] == "Pino Euro" else "USD"
            sums[cur] = round(sums.get(cur, 0) + r["Amount"], 2)
        page_totals = {k[len("total-"):]: float(v.replace(",", "")) for k, v in figs.items() if k.startswith("total-")}
        return int(figs["rows"]) == len(rows_) and page_totals == sums, f"{figs} vs {sums}"

    # ---------------------------------------------------------------- Company
    rfmd_page = client.get(f"/bank/export?{AUG}&owner={RFMD}").data.decode()
    check("2. Company report: READY, other owners' problems do not block it",
          "EXPORT READY" in rfmd_page and 'data-group="to_reconcile"' not in rfmd_page)
    rfmd_file = download(RFMD)
    header, rfmd_rows, rfmd_report = open_file(rfmd_file)
    check("2b. the file name names owner, year and month",
          'filename="RfBank_2026_08_RFMD.xlsx"' in rfmd_file.headers.get("Content-Disposition", ""))
    check("2c. the accountant's 12 columns stay first, then For Whom, Owner, Owner Type",
          tuple(header) == export_service.EXPORT_COLUMNS
          and tuple(header[:12]) == export_service.KERMALI_COLUMNS)
    check("2d. only RFMD's movements of August, each once — the card through its settlement account",
          sorted(r["Description"] for r in rfmd_rows)
          == ["RFMD CARD PURCHASE", "RFMD SUPPLIER ONE", "RFMD SUPPLIER TWO"])
    check("2e. RFMD rows carry the Company and the owner",
          all(r["Company"] == "RF Mount Dora, LLC" and r["Owner"] == "RFMD" and r["Owner Type"] == "Company"
              for r in rfmd_rows))
    check("2f. the Report sheet states owner, year and month",
          rfmd_report.get("Owner") == "RFMD" and rfmd_report.get("Year") == 2026 and rfmd_report.get("Month") == 8)
    ok, detail = matches(rfmd_page, rfmd_rows)
    check("8. Company: Check and file agree (row count, totals per currency)", ok, detail)

    # ----------------------------------------------------------------- Person
    pino_page = client.get(f"/bank/export?{AUG}&owner={PINO}").data.decode()
    check("3. Person report: READY — a personal Source never blocks", "EXPORT READY" in pino_page)
    pino_file = download(PINO)
    _, pino_rows, pino_report = open_file(pino_file)
    by_desc = {r["Description"]: r for r in pino_rows}
    check("3b. Pino's movements only: personal account, euro account, the card after it moved to Pino",
          sorted(by_desc) == ["MOVING CARD AFTER SWITCH", "PINO EURO PAYMENT", "PINO PAYS FOR MD"])
    check("4. derived card ownership by date: the same card counts for Angeli before Aug 15, for Pino after",
          "MOVING CARD BEFORE SWITCH" not in by_desc and "MOVING CARD AFTER SWITCH" in by_desc)
    check("5. a payment from Pino's account for MD stays in Pino's report, For Whom = RFMD",
          by_desc["PINO PAYS FOR MD"]["For Whom"] == "RFMD" and by_desc["PINO PAYS FOR MD"]["Owner"] == "Pino"
          and "PINO PAYS FOR MD" not in {r["Description"] for r in rfmd_rows})
    check("3c. a person needs no LLC: Company is empty, Owner Type is Person",
          all(r["Company"] is None and r["Owner Type"] == "Person" for r in pino_rows))
    ok, detail = matches(pino_page, pino_rows)
    check("8b. Person: Check and file agree, totals per currency (USD and EUR)",
          ok and "data-figure=\"total-EUR\"" in pino_page and "data-figure=\"total-USD\"" in pino_page, detail)
    check("8c. the Report sheet carries the totals per currency",
          pino_report.get("Total EUR") == -50.0 and pino_report.get("Total USD") == round(-123.45 - 8.0, 2))

    # ------------------------------------------------------ Another owner blocked
    angeli_page = client.get(f"/bank/export?{AUG}&owner={ANGELI}").data.decode()
    check("6. the owner with an open row is NOT READY, and its card row before the switch is in its scope",
          "EXPORT NOT READY" in angeli_page and figures(angeli_page).get("total") == "2")
    angeli_try = download(ANGELI)
    check("6b. its Download is refused", "attachment" not in angeli_try.headers.get("Content-Disposition", ""))

    # ------------------------------------------------------------ Unattributed
    check("7. Sources with no owner are listed separately, on every owner's page, without blocking",
          'data-group="unattributed"' in rfmd_page and "Orphan Account" in rfmd_page
          and "Card Without Settlement" in rfmd_page and "EXPORT READY" in rfmd_page)
    all_rows = len(rfmd_rows) + len(pino_rows) + 2           # Angeli's two rows
    with SessionFactory() as s:
        in_scope = len(export_service.in_scope_transactions(s, 2026, 8))
    check("7b. no movement duplicated or lost: owners' rows + unattributed = the month",
          all_rows + 2 == in_scope, f"{all_rows}+2 vs {in_scope}")

    # ---------------------------------------------------------------- Empty month
    september = client.get(f"/bank/export?year=2026&month=9&owner={RFMD}").data.decode()
    check("9. a month without movements says so and is not READY",
          "No transactions for this owner and month" in september and "EXPORT READY" not in september
          and 'id="export-download"' not in september)
    sept_post = client.post("/bank/export", data={"year": "2026", "month": "9", "owner": RFMD, "csrf_token": csrf})
    check("9b. and it produces no file", "attachment" not in sept_post.headers.get("Content-Disposition", ""))

    # ----------------------------------------------------------- Source warning
    sources = client.get("/bank/sources").data.decode()
    check("10. the Source page no longer warns 'Owned by a person'",
          "Owned by a person" not in sources and "requires a Company" not in sources)
    check("10b. a Source with no owner is still flagged, once", sources.count("No owner (Company or person)") == 1)

    # ------------------------------------------------------- Spinner and cookie
    check("11. the page keeps the wait cursor / spinner, and the download releases it",
          "rf-one-busy.js" in rfmd_page and "bank_export_done=t123" in rfmd_file.headers.get("Set-Cookie", ""))

    check("12. Check and Download changed nothing in the data", facts() == before)

    print()
    print(f"{len(passed)} passed, {len(failed)} failed.")
    return 1 if failed else 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    finally:
        try:
            from db import _engine
            _engine.dispose()
            os.remove(_TEST_DB_PATH)
        except Exception:  # noqa: BLE001 — best-effort cleanup of the throwaway file
            pass
