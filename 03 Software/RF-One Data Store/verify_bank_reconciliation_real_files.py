#!/usr/bin/env python
"""Read-only verification of the REAL Bank Reconciliation source files
against the figures recorded in
`01 Domains/Shared Domains/Administration/Bank Reconciliation/
BANK_RECONCILIATION_MANUAL_IMPORT_NORMALIZATION_001.md` §7.

This is NOT part of the permanent automated test suite: it depends on a
local, non-canonical directory of real bank exports that must never be
required for CI or for any other developer's checkout (spec: "I test
automatici permanenti non devono dipendere da questo percorso assoluto").
It never touches a database, never writes to the source directory, and
never modifies `RfBank.xlsx` or any CSV — every file is opened read-only.

Usage:
    python verify_bank_reconciliation_real_files.py <bank_folder>

<bank_folder> must contain a `Download/` subfolder with the CSV files and
an `RfBank.xlsx` file (i.e. the layout of the real, untracked `Bank/`
folder this task was told about). Can also be supplied via the
`BANK_RECON_REAL_FILES_DIR` environment variable. If neither is given, or
the path does not exist, this prints a message and exits 0 (skipped) —
its absence must never fail a normal test run.
"""

from __future__ import annotations

import os
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from rfone_data_store.bank_reconciliation import parsers, service


def _group_key(filename: str, account_hint: str | None) -> str:
    """For THIS DIAGNOSTIC SCRIPT ONLY — never for production import
    (spec §3.2 explicitly forbids resolving a real import by file name).
    Falls back to digits found in the file name only when the row itself
    carries no identifier, purely to group rows for this one-off,
    read-only cross-file comparison."""
    if account_hint:
        return account_hint
    digits = "".join(ch for ch in filename if ch.isdigit())[:4]
    return digits or filename


def _read_rfbank_history(xlsx_path: Path) -> set[tuple[date, str, int]]:
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    try:
        ws = wb["Bank"] if "Bank" in wb.sheetnames else wb[wb.sheetnames[0]]
        history: set[tuple[date, str, int]] = set()
        rows = ws.iter_rows(min_row=2, values_only=True)
        for row in rows:
            if len(row) < 4:
                continue
            raw_date, raw_description, raw_amount = row[1], row[2], row[3]
            if raw_date is None or raw_description is None or raw_amount is None:
                continue
            if isinstance(raw_date, datetime):
                d = raw_date.date()
            elif isinstance(raw_date, date):
                d = raw_date
            elif isinstance(raw_date, str):
                d = None
                for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d"):
                    try:
                        d = datetime.strptime(raw_date.strip(), fmt).date()
                        break
                    except ValueError:
                        continue
                if d is None:
                    continue
            else:
                continue
            try:
                amount_minor = int(round(float(raw_amount) * 100))
            except (TypeError, ValueError):
                continue
            history.add((d, service.normalize_description(str(raw_description)), amount_minor))
        return history
    finally:
        wb.close()


def main() -> int:
    folder = sys.argv[1] if len(sys.argv) > 1 else os.environ.get("BANK_RECON_REAL_FILES_DIR")
    if not folder:
        print("BANK_RECON_REAL_FILES_DIR not set and no path given — skipping (this is expected in CI).")
        return 0
    base = Path(folder)
    download_dir = base / "Download"
    rfbank_path = base / "RfBank.xlsx"
    if not download_dir.is_dir():
        print(f"'{download_dir}' does not exist — skipping (this is expected in CI).")
        return 0

    csv_files = sorted(download_dir.glob("*.csv"))
    print(f"Found {len(csv_files)} CSV file(s) in {download_dir}")

    all_rows: list[dict] = []  # one dict per parsed row, across every file
    per_file_row_count: dict[str, int] = {}

    for path in csv_files:
        data = path.read_bytes()
        parsed = parsers.parse_csv_bytes(data)
        per_file_row_count[path.name] = len(parsed.rows)
        for row in parsed.rows:
            all_rows.append({
                "file": path.name,
                "row": row,
                "group_key": _group_key(path.name, row.account_hint),
            })

    total_files = len(csv_files)
    total_rows = len(all_rows)
    missing_required_field = [
        r for r in all_rows
        if r["row"].posting_date is None or not r["row"].description or r["row"].amount_minor is None
    ]

    print(f"Total files: {total_files}")
    print(f"Total rows: {total_rows}")
    print(f"Rows missing date/description/amount: {len(missing_required_field)}")

    # -----------------------------------------------------------------
    # Certain overlaps (spec §7.2): the two Chase 1057 files; Chase 2270
    # fully contained in Chase 3144.
    # -----------------------------------------------------------------
    def file_signature_set(filename: str) -> set[tuple]:
        return {
            (r["group_key"], r["row"].posting_date, service.normalize_description(r["row"].description or ""),
             r["row"].amount_minor)
            for r in all_rows if r["file"] == filename
        }

    file_1057_with_card = "Chase1057_Activity_20260915.csv"
    file_1057_no_card = "Chase1057_Activity_20260915 (1).csv"
    file_2270 = "Chase2270_Activity_20260915.csv"
    file_3144 = "Chase3144_Activity_20260915.csv"

    present = {p.name for p in csv_files}
    overlap_1057 = 0
    if file_1057_with_card in present and file_1057_no_card in present:
        sig_a = {(d, desc, amt) for (_, d, desc, amt) in file_signature_set(file_1057_with_card)}
        sig_b = {(d, desc, amt) for (_, d, desc, amt) in file_signature_set(file_1057_no_card)}
        overlap_1057 = len(sig_a & sig_b)
        print(f"Chase 1057 pair: {per_file_row_count.get(file_1057_with_card)} vs "
              f"{per_file_row_count.get(file_1057_no_card)} rows, {overlap_1057} identical (date, description, amount) matches")
    else:
        print("Chase 1057 pair not found under the expected names — cannot verify this overlap.")

    overlap_2270_in_3144 = 0
    if file_2270 in present and file_3144 in present:
        sig_2270 = {(d, desc, amt) for (_, d, desc, amt) in file_signature_set(file_2270)}
        sig_3144 = {(d, desc, amt) for (_, d, desc, amt) in file_signature_set(file_3144)}
        overlap_2270_in_3144 = len(sig_2270 & sig_3144)
        print(f"Chase 2270 ({per_file_row_count.get(file_2270)} rows) rows also found in Chase 3144: {overlap_2270_in_3144}")
    else:
        print("Chase 2270 / Chase 3144 not found under the expected names — cannot verify this overlap.")

    certain_redundant = per_file_row_count.get(file_1057_no_card, 0) + per_file_row_count.get(file_2270, 0)
    print(f"Rows certainly redundant across downloads (entire redundant files): {certain_redundant}")

    remaining_rows = [
        r for r in all_rows if r["file"] not in (file_1057_no_card, file_2270)
    ]
    rows_after_exclusion = len(remaining_rows)
    print(f"Rows after excluding the two redundant downloads: {rows_after_exclusion}")

    # -----------------------------------------------------------------
    # Identical rows WITHIN the retained files (spec §7.3).
    # -----------------------------------------------------------------
    signature_groups: dict[tuple, list] = defaultdict(list)
    for r in remaining_rows:
        row = r["row"]
        sig = (r["group_key"], row.posting_date, service.normalize_description(row.description or ""), row.amount_minor)
        signature_groups[sig].append(r)
    excess_occurrences = sum(len(v) - 1 for v in signature_groups.values() if len(v) > 1)
    print(f"Excess identical occurrences within the retained files: {excess_occurrences}")
    for sig, group in signature_groups.items():
        if len(group) > 1:
            print(f"  - {len(group)}x {sig[1]} {sig[2]!r} {sig[3] / 100:.2f} (file(s): {sorted({g['file'] for g in group})})")

    # -----------------------------------------------------------------
    # Overlap with RfBank.xlsx history (spec §7.5).
    # -----------------------------------------------------------------
    if rfbank_path.exists():
        history = _read_rfbank_history(rfbank_path)
        matches = 0
        for r in remaining_rows:
            row = r["row"]
            if row.posting_date is None or row.amount_minor is None:
                continue
            key = (row.posting_date, service.normalize_description(row.description or ""), row.amount_minor)
            if key in history:
                matches += 1
        print(f"Rows also found in RfBank.xlsx history (date+description+amount match, account not compared): {matches}")
    else:
        print(f"'{rfbank_path}' not found — cannot verify overlap with RfBank history.")

    print()
    print("Expected (per BANK_RECONCILIATION_MANUAL_IMPORT_NORMALIZATION_001.md §7):")
    print("  14 files, 3174 rows, 0 rows missing date/description/amount,")
    print("  108 identical rows between the two Chase 1057 files,")
    print("  157 Chase 2270 rows already in Chase 3144,")
    print("  265 certainly redundant rows, 2909 rows after exclusion,")
    print("  6 excess identical occurrences, at least 841 matches with RfBank history.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
